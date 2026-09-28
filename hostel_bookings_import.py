# -*- coding: utf-8 -*-
"""Exely "bronlar hisoboti" Excel faylini o'qib, Xonalar moduli uchun
bronlar ro'yxatiga aylantiradi (mehmon F.I.Sh., kelish/ketish sanasi, xona
turi/raqami, holat, izoh, bron raqami).

Ustun nomlari til/versiyaga qarab farq qilgani uchun ustunlar pozitsiya
bo'yicha emas, kalit so'zlar bo'yicha aniqlanadi (exely_expense_import.py
bilan bir xil yondashuv). Har bir maydon uchun BIRINCHI mos ustun olinadi;
maydonlar ustuvorlik tartibida tekshiriladi, shuning uchun "Номер брони"
xona raqami emas, bron raqami bo'lib qoladi."""
import io
from datetime import date as date_cls, datetime

import openpyxl

# Ustuvorlik tartibi muhim: aniqroq maydonlar (ref, room_type) oldinda.
FIELD_HINTS = [
    ("ref", ["номер брони", "№ брони", "бронь №", "номер бронирования", "booking number", "booking no",
             "booking #", "booking id", "bron raqami", "bron №", "reservation number", "confirmation"]),
    ("room_type", ["тип номера", "категория номера", "категория", "room type", "xona turi", "тип комнаты"]),
    ("room", ["номер комнаты", "№ комнаты", "комната", "room number", "room no", "xona raqami", "xona",
              "room", "номер"]),
    ("guest", ["фио", "ф.и.о", "гость", "клиент", "имя гостя", "guest", "customer", "mehmon", "f.i.sh", "ism"]),
    ("arrival", ["заезд", "прибыт", "дата заезда", "arrival", "check-in", "check in", "checkin", "kelish"]),
    ("departure", ["выезд", "отъезд", "дата выезда", "departure", "check-out", "check out", "checkout", "ketish"]),
    ("status", ["статус", "состояние", "status", "holat"]),
    ("note", ["комментар", "примеч", "заметк", "comment", "note", "remark", "izoh"]),
]
LAST_HINTS = ["фамилия", "last name", "surname", "familiya"]
FIRST_HINTS = ["имя", "first name", "given name"]


def _norm(v):
    return " ".join(str(v or "").strip().lower().split())


def _classify(cell_text, taken):
    t = _norm(cell_text)
    if not t:
        return None
    for field, hints in FIELD_HINTS:
        if field in taken:
            continue
        if any(h in t for h in hints):
            return field
    return None


def _find_header(ws, max_scan=15):
    best = (None, None, 0)
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        taken = {}
        for col, val in enumerate(row):
            f = _classify(val, taken)
            if f:
                taken[f] = col
        names = {}
        for col, val in enumerate(row):
            t = _norm(val)
            if any(h == t or t.startswith(h) for h in LAST_HINTS) and "last" not in names:
                names["last"] = col
            elif any(h == t or t.startswith(h) for h in FIRST_HINTS) and "first" not in names:
                names["first"] = col
        score = len(taken) + (1 if ("last" in names or "first" in names) else 0)
        has_guest = "guest" in taken or ("last" in names and "first" in names)
        if has_guest and "arrival" in taken and score > best[2]:
            best = (idx, (taken, names), score)
    return best[0], best[1]


def _parse_date(v):
    if isinstance(v, (datetime, date_cls)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    if not s:
        return None
    for sep in (" ", ",", "T"):
        s = s.split(sep)[0]
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _text(v):
    return str(v).strip() if v not in (None, "") else ""


def parse_bookings_xlsx(file_bytes):
    """Qaytaradi: (rows, error, columns). rows — [{ref, guest, arrival,
    departure, room_type, room, status, note}]; error — 'col_not_found'
    yoki None; columns — aniqlangan {maydon: ustun_raqami} (oldindan
    ko'rishda foydalanuvchiga ko'rsatish uchun)."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active
    header_row, found = _find_header(ws)
    if header_row is None:
        return [], "col_not_found", {}
    taken, names = found

    rows = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v in (None, "") for v in values):
            continue

        def get(field):
            idx = taken.get(field)
            return values[idx] if idx is not None and idx < len(values) else None

        if "guest" in taken:
            guest = _text(get("guest"))
        else:
            parts = [_text(values[names[k]]) for k in ("last", "first") if k in names and names[k] < len(values)]
            guest = " ".join(p for p in parts if p)
        arrival = _parse_date(get("arrival"))
        if not guest or not arrival:
            continue
        rows.append({
            "ref": _text(get("ref")),
            "guest": guest,
            "arrival": arrival,
            "departure": _parse_date(get("departure")) or "",
            "room_type": _text(get("room_type")),
            "room": _text(get("room")),
            "status": _text(get("status")),
            "note": _text(get("note")),
        })
    columns = dict(taken)
    return rows, None, columns
