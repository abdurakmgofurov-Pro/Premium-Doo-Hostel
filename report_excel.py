# -*- coding: utf-8 -*-
"""Aggregate qilingan API ma'lumoti + qo'lda kiritilgan tranzaksiyalardan Excel hisobot yaratish."""
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from forma2 import build_forma2

HEADER_FILL = PatternFill(start_color="4B2E83", end_color="4B2E83", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def _style_header(ws, row=1, ncols=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


def build_excel_report(agg, manual, output_path):
    wb = openpyxl.Workbook()

    # --- Xulosa ---
    ws1 = wb.active
    ws1.title = "Xulosa"
    ws1["A1"] = "Premium Doo Hostel — Moliyaviy hisobot (Exely API)"
    ws1["A1"].font = TITLE_FONT
    ws1["A2"] = (
        f"Tayyorlangan: {datetime.now():%d.%m.%Y %H:%M} | "
        f"Jami bronlar: {agg['total_bookings']} (faol: {agg['active_count']}, "
        f"bekor qilingan: {agg['cancelled_count']}, {agg['cancellation_rate']}%)"
    )
    ws1["A2"].font = Font(italic=True)

    r = 4
    ws1.cell(row=r, column=1, value="Valyuta").font = BOLD
    ws1.cell(row=r, column=2, value="Bronlar soni").font = BOLD
    ws1.cell(row=r, column=3, value="Jami tushum (bron qiymati)").font = BOLD
    ws1.cell(row=r, column=4, value="Jami oldindan to'langan").font = BOLD
    _style_header(ws1, row=r, ncols=4)
    r += 1
    for cur in sorted(agg["revenue_by_currency"]):
        ws1.cell(row=r, column=1, value=cur)
        ws1.cell(row=r, column=2, value=agg["count_by_currency"].get(cur, 0))
        ws1.cell(row=r, column=3, value=round(agg["revenue_by_currency"].get(cur, 0), 2))
        ws1.cell(row=r, column=4, value=round(agg["prepaid_by_currency"].get(cur, 0), 2))
        r += 1
    r += 1
    ws1.cell(row=r, column=1, value="Qo'lda kiritilgan xarajatlar / kompaniyalar bilan hisob-kitob:").font = BOLD
    r += 1
    ws1.cell(row=r, column=1, value="Daromad (income)")
    ws1.cell(row=r, column=2, value=f"UZS: {manual['income']['UZS']:,.2f}")
    ws1.cell(row=r, column=3, value=f"USD: {manual['income']['USD']:,.2f}")
    r += 1
    ws1.cell(row=r, column=1, value="Xarajat (expense)")
    ws1.cell(row=r, column=2, value=f"UZS: {manual['expense']['UZS']:,.2f}")
    ws1.cell(row=r, column=3, value=f"USD: {manual['expense']['USD']:,.2f}")
    r += 1
    ws1.cell(row=r, column=1, value="To'lanmagan (unpaid) xarajat")
    ws1.cell(row=r, column=2, value=f"UZS: {manual['unpaid_expense']['UZS']:,.2f}")
    ws1.cell(row=r, column=3, value=f"USD: {manual['unpaid_expense']['USD']:,.2f}")
    _autofit(ws1, [40, 22, 26, 24])

    # --- Kanallar bo'yicha ---
    ws2 = wb.create_sheet("Kanallar bo'yicha")
    ws2.append(["Kanal", "Valyuta", "Bronlar soni", "Jami tushum"])
    _style_header(ws2, ncols=4)
    for name, c in sorted(agg["by_channel"].items(), key=lambda x: -sum(x[1]["revenue"].values())):
        for cur, val in c["revenue"].items():
            ws2.append([name, cur, c["count"], round(val, 2)])
    _autofit(ws2, [35, 10, 14, 18])

    # --- Oylar bo'yicha ---
    ws3 = wb.create_sheet("Oylar bo'yicha")
    ws3.append(["Oy (check-in)", "UZS bronlar", "UZS summa", "USD bronlar", "USD summa"])
    _style_header(ws3, ncols=5)
    for month in sorted(agg["by_month"]):
        ws3.append([
            month,
            agg["count_month"][month].get("UZS", 0), round(agg["by_month"][month].get("UZS", 0), 2),
            agg["count_month"][month].get("USD", 0), round(agg["by_month"][month].get("USD", 0), 2),
        ])
    _autofit(ws3, [16, 14, 16, 14, 16])

    # --- Barcha bronlar ---
    ws4 = wb.create_sheet("Barcha bronlar")
    ws4.append(["Bron raqami", "Holat", "Mehmon", "Valyuta", "Tushum", "Oldindan to'langan", "Kanal", "Yaratilgan", "Check-in"])
    _style_header(ws4, ncols=9)
    for row in agg["rows"]:
        ws4.append([
            row["number"], row["status"], row["guest"], row["currency"],
            round(row["revenue"], 2), round(row["prepaid"], 2), row["channel_name"],
            row["created"], row["arrival"],
        ])
    _autofit(ws4, [30, 12, 22, 9, 14, 18, 28, 20, 20])

    # --- Forma 2 (Moliyaviy natijalar to'g'risidagi hisobot) ---
    ws_f2 = wb.create_sheet("Forma 2")
    ws_f2["A1"] = "Forma 2 — Moliyaviy natijalar to'g'risidagi hisobot (ichki, boshqaruv hisoboti)"
    ws_f2["A1"].font = TITLE_FONT
    ws_f2.merge_cells("A1:C1")
    r = 3
    ws_f2.cell(row=r, column=1, value="Ko'rsatkich")
    ws_f2.cell(row=r, column=2, value="UZS")
    ws_f2.cell(row=r, column=3, value="USD")
    _style_header(ws_f2, row=r, ncols=3)
    for row in build_forma2(agg, manual):
        r += 1
        label = ("    " * row["indent"]) + row["label"]
        ws_f2.cell(row=r, column=1, value=label).font = BOLD if row["bold"] else Font()
        ws_f2.cell(row=r, column=2, value=round(row["uzs"], 2)).font = BOLD if row["bold"] else Font()
        ws_f2.cell(row=r, column=3, value=round(row["usd"], 2)).font = BOLD if row["bold"] else Font()
    r += 2
    f2_note = (
        "Bu — ichki (boshqaruv) hisobot, rasmiy tasdiqlangan buxgalteriya hujjati emas. "
        "010-qator Exely API orqali avtomatik olinadi; xarajat qatorlari \"Xarajatlar\" "
        "sahifasida qo'lda kiritilgan turkumlar bo'yicha guruhlangan. Rasmiy taqdim etishdan "
        "oldin litsenziyalangan buxgalter tomonidan tekshirilishi tavsiya etiladi. "
        "Forma 1 (Balans) va Forma 3 (Pul oqimlari) uchun aktivlar/majburiyatlar va kassa "
        "ma'lumotlari kerak — bular hozircha bu tizimda yig'ilmaydi."
    )
    ws_f2.cell(row=r, column=1, value=f2_note).alignment = Alignment(wrap_text=True)
    ws_f2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ws_f2.row_dimensions[r].height = 75
    _autofit(ws_f2, [50, 20, 20])

    # --- Xarajatlar / kompaniyalar bilan hisob-kitob ---
    ws5 = wb.create_sheet("Xarajatlar")
    ws5.append(["Sana", "Turi", "Kategoriya", "Forma 2 guruhi", "Kompaniya/kontragent", "Tavsif", "Summa", "Valyuta", "Holati"])
    _style_header(ws5, ncols=9)
    import db as _db
    for t in _db.list_transactions():
        group_label = _db.FORMA2_GROUP_LABELS.get(_db.CATEGORY_GROUP.get(t["category"]), "—") if t["type"] == "expense" else "—"
        ws5.append([t["date"], t["type"], t["category"], group_label, t["counterparty"], t["description"],
                    t["amount"], t["currency"], t["status"]])
    _autofit(ws5, [14, 10, 30, 34, 24, 34, 14, 10, 10])

    note = (
        "Eslatma: 'Tushum' — Exely'dagi bron qiymati (priceAfterTax), 'Oldindan to'langan' — "
        "guest tomonidan oldindan kiritilgan summa. Bekor qilingan (Cancelled) bronlar "
        "tushumga kiritilmagan. Xarajatlar/kompaniyalar bilan hisob-kitob qo'lda kiritiladi "
        "(dashboard'dagi \"Xarajatlar\" sahifasi orqali)."
    )
    ws1.cell(row=ws1.max_row + 2, column=1, value=note).alignment = Alignment(wrap_text=True)
    ws1.merge_cells(start_row=ws1.max_row, start_column=1, end_row=ws1.max_row, end_column=4)

    wb.save(output_path)


def build_cash_excel_report(rows, output_path):
    """Kassa/Bank (cash_transactions) ro'yxatini alohida Excel faylga
    eksport qiladi — joriy filtr (davr/turi/statya/kontragent) bilan
    chaqiriladi, shu ro'yxatning o'zi (allaqachon filtrlangan) yoziladi."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kassa-Bank"
    ws.append(["Sana", "Turi", "Hisob", "Statya nomi", "Kod (Cash Flow)", "Kontragent", "Tavsif", "Summa", "Valyuta"])
    _style_header(ws, ncols=9)
    for r in rows:
        ws.append([
            r["date"], r["type"], r["source"], r.get("note_label") or "", r["category"],
            r["counterparty"], r["description"], r["amount"], r["currency"],
        ])
    _autofit(ws, [14, 10, 10, 24, 30, 24, 34, 14, 10])
    wb.save(output_path)
