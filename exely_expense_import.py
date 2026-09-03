# -*- coding: utf-8 -*-
"""Exely extranet'ning "Финансовый учет -> Расход за период" bo'limidan
"Экспорт в XLSX" bilan yuklab olingan faylni o'qib, bizning Xarajatlar
(transactions) jadvaliga import qilish uchun tayyorlaydigan parser.

Exely'da bu ma'lumot uchun rasmiy API yo'q (faqat brauzer panelida ko'rinadi),
shuning uchun admin faylni qo'lda eksport qilib yuklaydi. Ustun nomlari
til/versiyaga qarab biroz farq qilishi mumkin bo'lgani uchun ustunlar
qattiq pozitsiya bo'yicha emas, kalit so'zlar bo'yicha aniqlanadi."""
import io
from datetime import date as date_cls, datetime

import openpyxl

HEADER_HINTS = {
    "date": ["дата", "date", "sana"],
    "category": ["статья", "category", "statya", "turkum"],
    "name": ["наименован", "name", "nomi"],
    "counterparty": ["фио", "плательщ", "получат", "counterpart", "kontragent"],
    "amount": ["сумма", "amount", "summa"],
    "payment_method": ["способ", "payment", "usul", "to'lov"],
}


def _norm(s):
    return str(s or "").strip().lower()


def _match_header(cell_text, hints):
    t = _norm(cell_text)
    return any(h in t for h in hints)


def _find_header_row(ws, max_scan=15):
    for row_idx in range(1, max_scan + 1):
        cells = [c.value for c in ws[row_idx]]
        texts = [_norm(c) for c in cells]
        has_date = any(_match_header(t, HEADER_HINTS["date"]) for t in texts)
        has_amount = any(_match_header(t, HEADER_HINTS["amount"]) for t in texts)
        if has_date and has_amount:
            col_map = {}
            for col_idx, val in enumerate(cells):
                for field, hints in HEADER_HINTS.items():
                    if field not in col_map and _match_header(val, hints):
                        col_map[field] = col_idx
            return row_idx, col_map
    return None, None


def _parse_date(v):
    if isinstance(v, (datetime, date_cls)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    if not s:
        return None
    s = s.split(" ")[0]
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _parse_amount(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_expense_xlsx(file_bytes):
    """Qaytaradi: (rows, error).
    rows — [{date, category_raw, name, counterparty, amount, currency, payment_method}]
    error — ustunlar aniqlanmasa xabar (str), aks holda None."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    header_row, col_map = _find_header_row(ws)
    if header_row is None or "amount" not in col_map or "date" not in col_map:
        return [], "col_not_found"

    rows = []
    for r in ws.iter_rows(min_row=header_row + 1):
        values = [c.value for c in r]
        if all(v in (None, "") for v in values):
            continue

        def raw(field):
            idx = col_map.get(field)
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        date_str = _parse_date(raw("date"))
        amount = _parse_amount(raw("amount"))
        if not date_str or amount is None or amount == 0:
            continue

        def text(field):
            v = raw(field)
            return str(v).strip() if v not in (None, "") else ""

        rows.append({
            "date": date_str,
            "category_raw": text("category") or "-",
            "name": text("name"),
            "counterparty": text("counterparty"),
            "amount": amount,
            "currency": "UZS",
            "payment_method": text("payment_method"),
        })
    return rows, None
